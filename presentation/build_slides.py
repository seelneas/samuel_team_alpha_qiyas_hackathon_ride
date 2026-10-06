"""Build presentation/team_alpha_slides.pptx – the 5-slide deck (Deliverable F).

    python presentation/build_slides.py [TEAM_NAME]
"""
import sys
from pathlib import Path

import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Cm, Pt

ROOT = Path(__file__).resolve().parents[1]
FIG = ROOT / "figures"
TEAM = sys.argv[1] if len(sys.argv) > 1 else "alpha"
OUT = ROOT / "presentation" / f"team_{TEAM.lower()}_slides.pptx"

res = pd.read_csv(ROOT / "reports" / "model_results.csv")
final = res[res.model.str.contains("optimised")].iloc[0]
sn = res[res.model == "Seasonal-naive (all history)"].iloc[0]

BLUE, DARK, GREY = RGBColor(0x0B, 0x4F, 0x8A), RGBColor(0x22, 0x22, 0x22), RGBColor(0x66, 0x66, 0x66)
prs = Presentation()
prs.slide_width, prs.slide_height = Cm(33.867), Cm(19.05)
BLANK = prs.slide_layouts[6]


def slide(n, title, subtitle=None):
    s = prs.slides.add_slide(BLANK)
    bar = s.shapes.add_shape(1, 0, 0, prs.slide_width, Cm(0.35))
    bar.fill.solid(); bar.fill.fore_color.rgb = BLUE; bar.line.fill.background()
    tb = s.shapes.add_textbox(Cm(1.2), Cm(0.7), Cm(31), Cm(1.6)).text_frame
    p = tb.paragraphs[0]; p.text = title; p.font.size = Pt(28); p.font.bold = True; p.font.color.rgb = BLUE
    if subtitle:
        q = tb.add_paragraph(); q.text = subtitle; q.font.size = Pt(14); q.font.color.rgb = GREY
    ft = s.shapes.add_textbox(Cm(1.2), Cm(18.0), Cm(31), Cm(0.8)).text_frame.paragraphs[0]
    ft.text = f"Addis Ride Demand Forecasting · TEAM {TEAM.upper()}   |   {n}/5"; ft.font.size = Pt(10); ft.font.color.rgb = GREY
    return s


def bullets(s, items, x, y, w, h, size=15):
    tf = s.shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h)).text_frame
    tf.word_wrap = True
    for i, it in enumerate(items):
        lvl = 1 if it.startswith("  ") else 0
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = lvl
        # **bold** prefix support
        txt = it.strip()
        if txt.startswith("**") and "**" in txt[2:]:
            b, rest = txt[2:].split("**", 1)
            r1 = p.add_run(); r1.text = ("• " if lvl == 0 else "– ") + b; r1.font.bold = True
            r2 = p.add_run(); r2.text = rest
            runs = [r1, r2]
        else:
            r = p.add_run(); r.text = ("• " if lvl == 0 else "– ") + txt; runs = [r]
        for r in runs:
            r.font.size = Pt(size - 2 * lvl); r.font.color.rgb = DARK
        p.space_after = Pt(5)
    return tf


def pic(s, name, x, y, w=None, h=None):
    return s.shapes.add_picture(str(FIG / name), Cm(x), Cm(y), width=Cm(w) if w else None, height=Cm(h) if h else None)


def table(s, rows, x, y, widths, size=12):
    t = s.shapes.add_table(len(rows), len(rows[0]), Cm(x), Cm(y), Cm(sum(widths)), Cm(0.8 * len(rows))).table
    for j, w in enumerate(widths):
        t.columns[j].width = Cm(w)
    for i, r in enumerate(rows):
        for j, v in enumerate(r):
            c = t.cell(i, j); c.text = str(v)
            for p in c.text_frame.paragraphs:
                p.font.size = Pt(size); p.font.bold = i == 0
    return t


# ---------------------------------------------------------------- 1 problem + tables
s = slide(1, "The problem: how many rides, where, and when?",
          "Forecast hourly trips in 12 Addis Ababa zones for 1–14 November 2025 (4,032 zone-hours)")
bullets(s, [
    "**Why it matters: **drivers in the right zone at the right hour mean shorter waits and more completed trips",
    "**Trips** (85k rows, Jan–Oct): hourly requests per zone, plus fares, wait times and active drivers",
    "  messy: 3 date formats, 55 spellings of 12 zones, −1 codes, duplicates, an outage, a zone launched in March",
    "**Weather** (hourly, observed + November forecast): temperature, rain, humidity, wind",
    "  trap: recorded on UTC, not local time; −9999 rain codes; a July block in °F",
    "**Events** (165): football, concerts, conferences, road closures, holidays",
    "  messy: multi-zone and city-wide entries, missing or reversed end times, free-text crowd sizes, cancellations",
    "**Our goal: **one clean zone × hour table that joins all three, then a model that uses only what is known 14 days ahead",
], 1.2, 3.4, 31, 14, size=16)

# ---------------------------------------------------------------- 2 cleaning & integration
s = slide(2, "Cleaning & integration: one clock, one zone key, one row per zone-hour")
bullets(s, [
    "**The clock problem: **the weather file is on UTC. Raw temperature peaks at 12:00; after +3 h it peaks at 15:00 EAT, and rain–demand correlation jumps from 0.06 to 0.42",
    "**Join map: **zone-hour grid ← weather (many-to-one on the hour) ← events (interval join from 2 h before start to 2 h after end, per affected zone)",
    "**Key numbers (A4):**",
    "  38,574 timestamps parsed explicitly (no day/month swaps)",
    "  32,391 zone labels → 12 canonical zones",
    "  846 duplicates, 675 −1 codes, 123 impossible spikes removed",
    "  353 °F readings converted, 76 −9999 rain codes, 192 missing weather hours filled",
    "  42-hour city-wide outage and Ayat before launch excluded, not zero-filled",
    "  Weather match 100%; joins never change row counts; 13/13 integrity checks pass",
    "**Result: **master_train has 82,979 zone-hours, master_test has 4,032",
], 1.2, 3.0, 15.8, 14.5, size=13)
pic(s, "fig06_weather_timezone_check.png", 17.3, 3.0, h=6.6)
pic(s, "fig01_gaps_and_missingness.png", 17.3, 10.0, h=7.8)

# ---------------------------------------------------------------- 3 findings
s = slide(3, "What the data says: rain and events move demand — but not everywhere")
pic(s, "fig07_rain_effect.png", 1.0, 2.8, w=15.6)
pic(s, "fig08_event_study.png", 17.2, 2.8, w=15.6)
bullets(s, [
    "**Rain: **+13% in light rain, +30% moderate, +46% heavy — sub-linear. But **Merkato falls ~21%**: an open-air market loses shoppers when it rains",
], 1.0, 13.4, 15.6, 4.4, size=13)
bullets(s, [
    "**Events: **football lifts demand +75%, with the biggest surge in the **2 hours after the final whistle (+124%)**. Concerts +31%; road closures −24%. So event windows must extend past the end time",
], 17.2, 13.4, 15.6, 4.4, size=13)

# ---------------------------------------------------------------- 4 modelling
s = slide(4, "Modelling & evaluation: beat the baseline, honestly",
          "Chronological validation: train before 18 Oct, test on 18–31 Oct, plus 4 rolling 14-day folds")
pic(s, "fig10_model_comparison.png", 17.2, 3.4, w=15.6)
table(s, [["Model", "RMSE", "MAE"],
          ["Mean predictor", "27.54", "20.69"],
          ["Seasonal-naive (zone×weekday×hour)", f"{sn.rmse:.2f}", f"{sn.mae:.2f}"],
          ["Ridge / random forest", "11.29 / 10.31", "7.55 / 6.76"],
          ["LightGBM, tuned", "9.03", "6.06"],
          ["LightGBM optimised (final)", f"{final.rmse:.2f}", f"{final.mae:.2f}"]], 1.2, 3.6, [8.6, 3.4, 3.4], size=12)
bullets(s, [
    "**Ablation: **calendar + lags 9.74 → + weather 9.21 → + events 9.64 → both **9.06**: both joins help",
    "**Rolling origin: **9.21 ± 0.71 vs 12.32 ± 1.83 for seasonal-naive, better in every fold",
    "**Leakage check: **adding active_drivers fakes an RMSE of 4.25, so it is excluded; only lags of 14 days or more are used",
    "**Optimisation: **Poisson loss, 2–5-week lags and an 8-week zone×weekday×hour profile take RMSE from 9.03 to 8.47",
], 1.2, 9.0, 15.6, 8.6, size=12)

# ---------------------------------------------------------------- 5 errors, score, demo, next
s = slide(5, f"Error analysis, final score {final.rmse:.2f} RMSE, demo and next steps")
pic(s, "fig11_forecast_vs_actual.png", 17.2, 2.8, h=8.2)
bullets(s, [
    f"**Final score: **RMSE {final.rmse:.2f}, MAE {final.mae:.2f} trips per zone-hour (≈17.5% of demand). That is {100*(1-final.rmse/sn.rmse):.0f}% better than seasonal-naive and ≈ 4–5 drivers per zone-hour",
    "**Where it misses: **busy zones in absolute terms (Bole 11.0, Megenagna 10.9, Merkato 10.5 vs Ayat 5.9) but relative error is flat (17–20%); rush hours; the tallest post-event surges; one-off holidays (Meskel)",
    "**Demo: **pick a zone and a date → 24-h forecast, peak hour, drivers needed, expected fares, and the weather/events looked up",
    "  run locally: streamlit run app/app.py   (live demo on our laptop)",
    "**Next: **holiday- and event-specific multipliers, quantile forecasts for staffing, training with noisy weather forecasts, weekly retraining and hosting",
], 1.2, 2.8, 15.6, 15, size=13)
pic(s, "fig12_feature_importance.png", 17.2, 11.2, h=6.6)

# keep every picture inside its column
for sl in prs.slides:
    for sh in sl.shapes:
        if sh.shape_type == 13 and sh.left + sh.width > Cm(33.2):
            r = (Cm(33.2) - sh.left) / sh.width
            sh.width, sh.height = int(sh.width * r), int(sh.height * r)

OUT.parent.mkdir(exist_ok=True)
prs.save(OUT)
print("written", OUT)
